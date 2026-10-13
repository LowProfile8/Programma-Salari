"""Schede salari in Excel, sul modello dei fogli «Salari 2026 - <azienda>»: una colonna per mese con tutte le voci della busta paga
(scheda del singolo dipendente) e il riepilogo di tutti i dipendenti (salari totali)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.buste.calcolo import MESI, buste_dipendente
from core.util import data_it

BLU = "07437A"
VOCI = [  # (etichetta, chiave in calc, tipo)  tipo: n = importo, t = totale (grassetto)
    ("Stipendio lordo", "stipendio", "n"), ("Stipendio lordo periodo di carenza", "carenza", "n"),
    ("Supplemento vacanze", "vac", "n"), ("Supplemento festività", "fest", "n"), ("Quota mensile 13ma", "q13", "n"),
    ("Ore supplementari", "ore_supp", "n"), ("Indennità Maternità / Paternità", "maternita", "n"), ("Bonus", "bonus", "n"),
    ("Salario Lordo AVS", "lordo_avs", "t"), ("Assegni Familiari", "assegni", "n"), ("Indennità Assicurazione", "ind_ass", "n"),
    ("Salario Lordo soggetto IF", "lordo_if", "t"),
    ("Trattenuta AVS", "avs", "n"), ("Trattenuta AD", "ad", "n"), ("Trattenuta IGM/IPG", "igm", "n"), ("Trattenuta LAINF", "lainf", "n"),
    ("Trattenuta LPP", "lpp", "n"), ("Trattenuta Imposta alla fonte", "imposta", "n"), ("Vitto e alloggio", "vitto", "n"),
    ("Contributo sindacale / altre trattenute", "_altre", "n"), ("Conguagli oneri sociali", "cong_tot", "n"), ("Totale deduzioni", "deduzioni", "t"),
    ("Rimborso lavaggio abiti", "lavaggio", "n"), ("Rimborsi diversi", "rimborsi", "n"), ("Arrotondamento", "arrot", "n"),
    ("Salario netto in CHF", "netto", "t"),
]
_FILL = PatternFill("solid", fgColor=BLU)
_CHIARO = PatternFill("solid", fgColor="E8F0FA")
_LATO = Side(style="thin", color="B8C4D6")
_BORDO = Border(left=_LATO, right=_LATO, top=_LATO, bottom=_LATO)
_NUM = '#,##0.00;[Red]-#,##0.00;"-"'


def _v(calc: dict, k: str) -> float:
    if k == "_altre":
        return float(calc.get("altra") or 0) + float(calc.get("contributo") or 0)
    return float(calc.get(k) or 0)


def _testa(c, testo: str, larg: bool = False) -> None:
    c.value = testo
    c.fill, c.font = _FILL, Font(bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="center" if not larg else "left", vertical="center", wrap_text=True)
    c.border = _BORDO


def scheda_salari_xlsx(az: dict, dip: dict, anno: int) -> bytes:
    """Scheda salari del dipendente per un anno: voci in riga, mesi in colonna, totale dell'anno."""
    buste = {b["mese"]: b["calc"] for b in buste_dipendente(dip) if b["anno"] == anno}
    wb = Workbook()
    ws = wb.active
    ws.title = f"Scheda salari {anno}"
    ws["B2"] = f"Scheda salari {anno}"
    ws["B2"].font = Font(bold=True, size=16, color=BLU)
    info = [("Datore di lavoro", az.get("ragione_sociale") or ""),
            ("Sede", ", ".join(x for x in (az.get("sede_via"), " ".join(y for y in (az.get("sede_npa"), az.get("sede_localita")) if y)) if x)),
            ("Dipendente", f"{dip.get('cognome', '')} {dip.get('nome', '')}".strip()),
            ("N. AVS", dip.get("numero_avs") or ""), ("Data di nascita", data_it(dip.get("data_nascita")) if dip.get("data_nascita") else ""),
            ("Data di entrata", data_it(dip.get("data_assunzione")) if dip.get("data_assunzione") else ""),
            ("Impiego", dip.get("funzione") or "")]
    for i, (e, v) in enumerate(info):
        ws.cell(4 + i, 2, e).font = Font(bold=True)
        ws.cell(4 + i, 3, v)
    r0 = 4 + len(info) + 1
    _testa(ws.cell(r0, 2), "Conteggio paga", True)
    for m in range(1, 13):
        _testa(ws.cell(r0, 2 + m), MESI[m - 1])
    _testa(ws.cell(r0, 15), f"Totale {anno}")
    for i, (et, k, tipo) in enumerate(VOCI):
        r = r0 + 1 + i
        c = ws.cell(r, 2, et)
        c.border = _BORDO
        if tipo == "t":
            c.font = Font(bold=True)
        tot = 0.0
        for m in range(1, 13):
            x = ws.cell(r, 2 + m)
            x.border = _BORDO
            x.number_format = _NUM
            if m in buste:
                x.value = _v(buste[m], k)
                tot += x.value
            if tipo == "t":
                x.font, x.fill = Font(bold=True), _CHIARO
        t = ws.cell(r, 15, tot)
        t.number_format, t.border, t.font = _NUM, _BORDO, Font(bold=True)
        if tipo == "t":
            c.fill = _CHIARO
            t.fill = _CHIARO
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 38
    for col in range(3, 16):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.freeze_panes = ws.cell(r0 + 1, 3)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def salari_totali_xlsx(az: dict, anno: int) -> bytes:
    """Riepilogo di tutti i dipendenti dell'azienda: un foglio per voce (dipendenti in riga, mesi in colonna, totale per persona e per mese)."""
    wb = Workbook()
    wb.remove(wb.active)
    dips = sorted(az.get("dipendenti", []), key=lambda d: (d.get("cognome", "").casefold(), d.get("nome", "").casefold()))
    dati = []
    for d in dips:
        b = {x["mese"]: x["calc"] for x in buste_dipendente(d) if x["anno"] == anno}
        if b:
            dati.append((d, b))
    for et, k, _ in VOCI:
        ws = wb.create_sheet(et[:31].replace("/", "-"))
        ws["B2"] = f"{az.get('ragione_sociale', '')} — {et} {anno}"
        ws["B2"].font = Font(bold=True, size=14, color=BLU)
        for j, t in enumerate(["Nome e cognome", "N. AVS", "Data di nascita"], start=2):
            _testa(ws.cell(4, j), t, True)
        for m in range(1, 13):
            _testa(ws.cell(4, 4 + m), MESI[m - 1])
        _testa(ws.cell(4, 17), f"Totale {anno}")
        for i, (d, b) in enumerate(dati):
            r = 5 + i
            ws.cell(r, 2, f"{d.get('cognome', '')} {d.get('nome', '')}".strip())
            ws.cell(r, 3, d.get("numero_avs") or "")
            ws.cell(r, 4, data_it(d.get("data_nascita")) if d.get("data_nascita") else "")
            tot = 0.0
            for m in range(1, 13):
                c = ws.cell(r, 4 + m)
                if m in b:
                    c.value = _v(b[m], k)
                    tot += c.value
                c.number_format = _NUM
            t = ws.cell(r, 17, tot)
            t.number_format, t.font = _NUM, Font(bold=True)
        rt = 5 + len(dati)
        ws.cell(rt, 2, "Totale per mese").font = Font(bold=True)
        for col in range(5, 18):
            L = get_column_letter(col)
            c = ws.cell(rt, col, f"=SUM({L}5:{L}{rt - 1})" if dati else 0)
            c.number_format, c.font, c.fill = _NUM, Font(bold=True), _CHIARO
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 14
        for col in range(5, 18):
            ws.column_dimensions[get_column_letter(col)].width = 13
        ws.freeze_panes = "E5"
    if not wb.sheetnames:
        wb.create_sheet("Salari totali")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def tabella_totali(az: dict, anno: int, chiave: str) -> list[dict]:
    """Righe per la tabella a schermo: un dipendente per riga, mesi in colonna e totale."""
    dips = sorted(az.get("dipendenti", []), key=lambda d: (d.get("cognome", "").casefold(), d.get("nome", "").casefold()))
    righe = []
    for d in dips:
        b = {x["mese"]: x["calc"] for x in buste_dipendente(d) if x["anno"] == anno}
        if not b:
            continue
        riga = {"Dipendente": f"{d.get('cognome', '')} {d.get('nome', '')}".strip()}
        tot = 0.0
        for m in range(1, 13):
            v = _v(b[m], chiave) if m in b else None
            riga[MESI[m - 1][:3]] = v
            tot += v or 0.0
        riga["Totale"] = tot
        righe.append(riga)
    if righe:
        tot = {"Dipendente": "Totale"}
        for col in list(righe[0])[1:]:
            tot[col] = sum(r[col] or 0.0 for r in righe)
        righe.append(tot)
    return righe


def tabella_generale(az: dict, anno: int, mese: int | None = None) -> list[dict]:
    """Tabella generale: un dipendente per riga e tutte le voci della busta paga una accanto all'altra (somma dell'anno o di un solo mese)."""
    dips = sorted(az.get("dipendenti", []), key=lambda d: (d.get("cognome", "").casefold(), d.get("nome", "").casefold()))
    righe = []
    for d in dips:
        bs = [x["calc"] for x in buste_dipendente(d) if x["anno"] == anno and (mese is None or x["mese"] == mese)]
        if not bs:
            continue
        riga = {"Dipendente": f"{d.get('cognome', '')} {d.get('nome', '')}".strip(), "Buste": len(bs)}
        for et, k, _ in VOCI:
            riga[et] = sum(_v(c, k) for c in bs)
        righe.append(riga)
    if righe:
        tot = {"Dipendente": "Totale", "Buste": sum(r["Buste"] for r in righe)}
        for et, _, _ in VOCI:
            tot[et] = sum(r[et] for r in righe)
        righe.append(tot)
    return righe


# ── riconciliazione salariale (registrazioni di contabilità: conti 5010-5140) ──
def _somme(az: dict, anno: int, mese: int | None) -> dict:
    t = {k: 0.0 for k in ("avs_ad", "igm", "lainf", "lpp", "imposta", "paritetica", "viaggio", "altri", "trasferta", "ass")}
    for d in az.get("dipendenti", []):
        for b in buste_dipendente(d):
            if b["anno"] != anno or (mese is not None and b["mese"] != mese):
                continue
            c = b["calc"] or {}
            cg = c.get("cong") or {}
            g = lambda k: float(c.get(k) or 0.0)       # noqa: E731
            t["avs_ad"] += g("avs") + g("ad")
            t["igm"] += g("igm") + float(cg.get("igm") or 0)
            t["lainf"] += g("lainf") + float(cg.get("lainf") or 0)
            t["lpp"] += g("lpp") + float(cg.get("lpp") or 0)
            t["imposta"] += g("imposta") + float(cg.get("imposta") or 0)
            t["paritetica"] += g("contributo")
            nome_altra = str((b.get("input") or {}).get("altra_nome") or "").lower()
            if "viagg" in nome_altra:
                t["viaggio"] += g("altra")
                t["altri"] += g("vitto") - g("arrot")
            else:
                t["altri"] += g("vitto") + g("altra") - g("arrot")
            t["trasferta"] += g("lavaggio") + g("rimborsi")
            t["ass"] += g("ind_ass")
    return t


_RIGHE_RICONC = [  # (descrizione, conto dare, conto avere, chiave)
    ("Deduzione contributi AVS, AI, AD", 5010, 5020, "avs_ad"), ("Deduzione contributi IGM", 5010, 5040, "igm"),
    ("Deduzione contributi LAINF", 5010, 5050, "lainf"), ("Deduzione contributi LPP", 5010, 5060, "lpp"),
    ("Deduzione contributi Imposta alla Fonte", 5010, 5070, "imposta"),
    ("Deduzione contributi commissione parietetica", 5010, 5100, "paritetica"),
    ("Deduzione contributi spese di viaggio", 5010, 5120, "viaggio"),
    ("Altri costi del personale e arrotondamenti", 5010, 5110, "altri"),
    ("Indennità di trasferta", 5140, 5010, "trasferta"), ("Indennità Assicurazione", 5130, 5010, "ass"),
]


def riconciliazione_xlsx(az: dict, anno: int) -> bytes:
    """Riconciliazione salariale: le registrazioni da fare in contabilità (totali dell'anno al 31.12 e, nel secondo foglio, mese per mese)."""
    import datetime as dt
    import calendar
    wb = Workbook()
    for titolo, mesi in ((f"Riconciliazione {anno}", [None]), ("Per mese", list(range(1, 13)))):
        ws = wb.active if titolo.startswith("Riconc") else wb.create_sheet(titolo)
        ws.title = titolo
        for j, e in enumerate(["Data", "Riferimento", "Descrizione", "Conto dare", "Conto avere", "Importo", "Valuta", "Cambio", "Importo CHF"], start=1):
            _testa(ws.cell(1, j), e, j == 3)
        r = 2
        for m in mesi:
            data = dt.date(anno, 12, 31) if m is None else dt.date(anno, m, calendar.monthrange(anno, m)[1])
            somme = _somme(az, anno, m)
            for desc, dare, avere, k in _RIGHE_RICONC:
                ws.cell(r, 1, data).number_format = "DD.MM.YYYY"
                ws.cell(r, 3, desc)
                ws.cell(r, 4, dare)
                ws.cell(r, 5, avere)
                ws.cell(r, 6, round(somme[k], 2)).number_format = "#,##0.00"
                ws.cell(r, 7, "CHF")
                ws.cell(r, 8, 1.0).number_format = "0.000000"
                ws.cell(r, 9, f"=F{r}*H{r}").number_format = "#,##0.00"
                r += 1
        for col, w in zip("ABCDEFGHI", (13, 12, 48, 12, 12, 16, 8, 11, 16)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
